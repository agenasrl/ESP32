#!/usr/bin/env python3
"""Convert a pick-and-place export (.txt/.csv/.xlsx/.xls) into CPL_mm.xlsx.

Output columns follow the BV2301A0_CPL_mm template:
    Designator | Mid X | Mid Y | Layer | Rotation
Coordinates are written in mm, Layer as TOP / BOT.

Supported inputs (auto-detected):
  * Altium "Pick and Place Locations" text (fixed-width, quoted descriptions,
    possibly spanning several lines)
  * Cadence/Allegro place_components text ("!"-separated, mirror column)
  * Generic CSV / TSV / whitespace tables with a header row
  * Excel sheets (.xlsx/.xlsm, .xls) with a header row in the first 30 rows
"""
import argparse
import csv
import os
import re
import shlex
import shutil
import subprocess
import sys
import tempfile

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
TEMPLATE = os.path.join(HERE, "..", "assets", "CPL_template.xlsx")

# Normalised header names, best match first.
FIELDS = {
    "designator": ["designator", "refdes", "reference", "ref", "refdesignator", "partreference", "component", "part"],
    "x": ["midx", "centerx", "symbolx", "centroidx", "posx", "locationx", "placex", "x", "refx"],
    "y": ["midy", "centery", "symboly", "centroidy", "posy", "locationy", "placey", "y", "refy"],
    "layer": ["layer", "side", "tb", "placementlayer", "boardside"],
    "mirror": ["mirror", "mirrored", "flip"],
    "rotation": ["rotation", "rot", "angle", "orientation", "symbolrotation"],
}


def norm(h):
    h = str(h or "").lower()
    h = re.sub(r"\((mm|mil|mils|inch|in)\)", "", h)
    return re.sub(r"[^a-z0-9]", "", h)


def unit_of(h):
    m = re.search(r"\((mm|mil|mils|inch|in)\)", str(h or "").lower())
    return m.group(1) if m else None


def map_columns(header):
    keys = [norm(h) for h in header]
    cols = {}
    for field, names in FIELDS.items():
        for name in names:
            if name in keys:
                cols[field] = keys.index(name)
                break
    return cols


def is_header(cells):
    c = map_columns(cells)
    return "designator" in c and "x" in c and "y" in c


def to_mm(val, default_unit):
    s = str(val).strip().strip('"').lower().replace(",", ".")
    unit = default_unit
    m = re.match(r"^(-?[\d.]+)\s*(mm|mil|mils|in|inch)?$", s)
    if not m:
        raise ValueError(f"not a number: {val!r}")
    if m.group(2):
        unit = m.group(2)
    v = float(m.group(1))
    if unit in ("mil", "mils"):
        v *= 0.0254
    elif unit in ("in", "inch"):
        v *= 25.4
    return round(v, 4)


def layer_of(row, cols):
    if "layer" in cols:
        v = str(row[cols["layer"]]).strip().lower()
        if v.startswith("b") or "bot" in v:
            return "BOT"
        if v.startswith("t") or "top" in v:
            return "TOP"
    if "mirror" in cols:
        v = str(row[cols["mirror"]] or "").strip().lower()
        return "BOT" if v in ("m", "y", "yes", "true", "1", "mirror", "mirrored") else "TOP"
    return "TOP"


def rotation_of(row, cols):
    if "rotation" not in cols:
        return None
    s = str(row[cols["rotation"]]).strip()
    if not s:
        return None
    v = float(re.sub(r"[^\d.\-]", "", s))
    return int(v) if v == int(v) else v


def split_quoted(line, delim):
    if delim == "!":
        return [p.strip() for p in line.lstrip("#").split("!")]
    if delim in (",", "\t", ";"):
        return [p.strip() for p in next(csv.reader([line], delimiter=delim))]
    lex = shlex.shlex(line, posix=True)
    lex.whitespace_split = True
    lex.quotes = '"'
    lex.escape = ""
    lex.commenters = ""
    return list(lex)


def read_text(path):
    raw = open(path, "rb").read()
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    lines = text.replace("\r\n", "\n").replace("\r", "\n").split("\n")

    file_unit = "mm"
    m = re.search(r"units?\s*(used)?\s*[:=]\s*(millimeters?|mm|mils?|inch(es)?|in)\b", text, re.I)
    if m:
        u = m.group(2).lower()
        file_unit = "mm" if u.startswith("milli") else "inch" if u.startswith("in") else u

    # find the header line and its delimiter
    for i, line in enumerate(lines):
        for delim in ("!", "\t", ",", ";", None):
            if delim and delim not in line:
                continue
            try:
                cells = split_quoted(line, delim)
            except ValueError:
                continue
            if is_header(cells):
                header, start, sep = cells, i + 1, delim
                break
        else:
            continue
        break
    else:
        sys.exit("No header row with designator / X / Y columns found.")

    rows, buf = [], ""
    for line in lines[start:]:
        if sep == "!" and line.lstrip().startswith("#"):
            continue
        buf = f"{buf} {line}" if buf else line
        if sep is None and buf.count('"') % 2:
            continue  # description continues on the next line
        if buf.strip():
            rows.append(split_quoted(buf, sep))
        buf = ""
    return header, rows, file_unit


def read_excel(path):
    if path.lower().endswith(".xls"):
        try:
            import xlrd
            sh = xlrd.open_workbook(path).sheet_by_index(0)
            data = [sh.row_values(r) for r in range(sh.nrows)]
        except ImportError:
            office = shutil.which("soffice") or shutil.which("libreoffice")
            if not office:
                sys.exit("Reading .xls needs xlrd (pip install xlrd) or LibreOffice.")
            tmp = tempfile.mkdtemp()
            subprocess.run([office, "--headless", "--convert-to", "xlsx", "--outdir", tmp, path],
                           check=True, capture_output=True)
            path = os.path.join(tmp, os.path.splitext(os.path.basename(path))[0] + ".xlsx")
            return read_excel(path)
    else:
        ws = openpyxl.load_workbook(path, data_only=True).active
        data = [list(r) for r in ws.iter_rows(values_only=True)]
    for i, r in enumerate(data[:30]):
        cells = ["" if c is None else str(c) for c in r]
        if is_header(cells):
            body = [["" if c is None else c for c in row] for row in data[i + 1:]]
            return cells, body, "mm"
    sys.exit("No header row with designator / X / Y columns found in the first 30 rows.")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("input")
    ap.add_argument("-o", "--output", default="CPL_mm.xlsx")
    ap.add_argument("--force", action="store_true", help="overwrite an existing output file")
    a = ap.parse_args()

    if os.path.exists(a.output) and not a.force:
        sys.exit(f"{a.output} already exists; pass --force to overwrite.")

    ext = os.path.splitext(a.input)[1].lower()
    header, rows, file_unit = read_excel(a.input) if ext in (".xlsx", ".xlsm", ".xls") else read_text(a.input)
    cols = map_columns(header)
    ux = unit_of(header[cols["x"]]) or file_unit
    uy = unit_of(header[cols["y"]]) or file_unit

    out, skipped = [], []
    for r in rows:
        need = max(cols.values())
        if len(r) <= need or not str(r[cols["designator"]]).strip():
            if any(str(c).strip() for c in r):
                skipped.append(r)
            continue
        try:
            out.append((str(r[cols["designator"]]).strip(),
                        to_mm(r[cols["x"]], ux), to_mm(r[cols["y"]], uy),
                        layer_of(r, cols), rotation_of(r, cols)))
        except ValueError:
            skipped.append(r)

    wb = openpyxl.load_workbook(TEMPLATE)
    ws = wb.active
    for i, row in enumerate(out, start=2):
        for j, v in enumerate(row, start=1):
            c = ws.cell(i, j)
            c.value = v
            if j in (2, 3):
                c.number_format = "0.0000"
            elif j == 5:
                c.number_format = "0"
    wb.save(a.output)

    used = {f: header[i] for f, i in cols.items()}
    dups = sorted({d for d, *_ in out if [o[0] for o in out].count(d) > 1})
    print(f"wrote {a.output}: {len(out)} parts "
          f"({sum(o[3] == 'TOP' for o in out)} TOP, {sum(o[3] == 'BOT' for o in out)} BOT)")
    print(f"columns used: {used}; source units X={ux} Y={uy}")
    if "layer" not in cols and "mirror" not in cols:
        print("WARNING: no layer/side/mirror column - every part set to TOP")
    if "rotation" not in cols:
        print("WARNING: no rotation column - Rotation left empty")
    if dups:
        print(f"WARNING: duplicate designators: {', '.join(dups)}")
    for r in skipped:
        print(f"skipped row: {r}")


if __name__ == "__main__":
    main()
